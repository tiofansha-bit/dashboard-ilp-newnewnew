"""PWS ILP MELATI - Backend API
Digitalisasi Ceklis Kunjungan Rumah Kader - UPT Puskesmas Melati.
"""
from dotenv import load_dotenv
from pathlib import Path
import os
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from typing import List, Optional, Any, Dict
from datetime import datetime, timezone, timedelta, date
import uuid, logging, json, io, bcrypt, jwt, re
from collections import defaultdict, Counter

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("pws")

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALG = "HS256"

app = FastAPI(title="PWS ILP MELATI")
api = APIRouter(prefix="/api")

def now_utc():
    return datetime.now(timezone.utc)

def iso(dt=None):
    return (dt or now_utc()).isoformat()

def new_id():
    return str(uuid.uuid4())

# ---------------- Auth helpers ----------------
def hash_password(p: str) -> str:
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()

def verify_password(p: str, h: str) -> bool:
    try:
        return bcrypt.checkpw(p.encode(), h.encode())
    except Exception:
        return False

def create_token(user, kind="access"):
    exp = now_utc() + (timedelta(minutes=720) if kind == "access" else timedelta(days=7))
    return jwt.encode({"sub": user["id"], "role": user["role"], "type": kind, "exp": exp}, JWT_SECRET, algorithm=JWT_ALG)

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        h = request.headers.get("Authorization", "")
        if h.startswith("Bearer "):
            token = h[7:]
    if not token:
        raise HTTPException(401, "Belum masuk / sesi berakhir")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Sesi berakhir")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Token tidak valid")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(401, "Pengguna tidak ditemukan")
    return user

async def require_admin(user: dict = Depends(get_current_user)):
    if user["role"] != "admin":
        raise HTTPException(403, "Hanya untuk petugas/admin")
    return user

# ---------------- Audit ----------------
async def audit(user, aksi, entitas, entitas_id, detail=""):
    await db.audit_logs.insert_one({
        "id": new_id(), "user_id": user.get("id") if user else None,
        "user_nama": user.get("nama") if user else "system", "aksi": aksi,
        "entitas": entitas, "entitas_id": entitas_id, "detail": detail, "waktu": iso()})

# ---------------- Models ----------------
class LoginIn(BaseModel):
    username: str
    password: str

class KeluargaIn(BaseModel):
    nama_kk: str
    jumlah_anggota: Optional[int] = None
    alamat: str = ""
    rt: str = ""
    rw: str = ""
    kelurahan: str
    posyandu: str = ""
    no_hp: str = ""
    catatan_lokasi: str = ""
    gps: Optional[Dict[str, float]] = None
    punya_jkn: Optional[bool] = None
    air_bersih: Optional[bool] = None
    jamban: Optional[bool] = None
    ventilasi: Optional[bool] = None
    ada_gangguan_jiwa: Optional[bool] = None
    ada_tbc: Optional[bool] = None
    ada_hipertensi: Optional[bool] = None
    ada_dm: Optional[bool] = None

class AnggotaIn(BaseModel):
    keluarga_id: str
    nama: str
    nik: str = ""
    tempat_lahir: str = ""
    tanggal_lahir: Optional[str] = None
    jenis_kelamin: str = ""
    hubungan_kk: str = ""
    status_kawin: str = ""
    pendidikan: str = ""
    pekerjaan: str = ""
    no_hp: str = ""
    kelompok_override: Optional[str] = None
    kondisi_hamil: bool = False
    kondisi_nifas: bool = False

class KunjunganIn(BaseModel):
    keluarga_id: str
    anggota_ids: List[str] = []
    tanggal: Optional[str] = None
    catatan: str = ""
    gps: Optional[Dict[str, float]] = None
    answers: Dict[str, Dict[str, Any]] = {}   # {anggota_id: {kode: value}}
    edukasi: Dict[str, List[str]] = {}        # {anggota_id: [materi]}
    status: str = "draf"                       # draf | selesai
    reminder_confirmed: bool = False

class CaseUpdateIn(BaseModel):
    status: Optional[str] = None
    penanggung_jawab: Optional[str] = None
    jenis_tindak_lanjut: Optional[str] = None
    hasil: Optional[str] = None
    rencana: Optional[str] = None
    tempat: Optional[str] = None
    alasan: Optional[str] = None
    target_selesai: Optional[str] = None

# ---------------- Target group logic ----------------
def hitung_umur(tgl_lahir):
    if not tgl_lahir:
        return None
    try:
        d = datetime.fromisoformat(str(tgl_lahir)[:10]).date()
    except Exception:
        return None
    t = date.today()
    return t.year - d.year - ((t.month, t.day) < (d.month, d.day))

def umur_bulan(tgl_lahir):
    if not tgl_lahir:
        return None
    try:
        d = datetime.fromisoformat(str(tgl_lahir)[:10]).date()
    except Exception:
        return None
    t = date.today()
    return (t.year - d.year) * 12 + (t.month - d.month)

def tentukan_kelompok(a):
    if a.get("kelompok_override"):
        return a["kelompok_override"]
    if a.get("kondisi_hamil"):
        return "ibu_hamil"
    if a.get("kondisi_nifas"):
        return "nifas"
    mb = umur_bulan(a.get("tanggal_lahir"))
    th = hitung_umur(a.get("tanggal_lahir"))
    if mb is None:
        return "belum_ditentukan"
    if mb <= 6:
        return "bayi"
    if mb <= 71:
        return "balita"
    if th is not None and th < 18:
        return "remaja"
    if th is not None and th < 60:
        return "dewasa"
    return "lansia"

KELOMPOK_LABEL = {
    "ibu_hamil": "Ibu Hamil", "nifas": "Ibu Bersalin & Nifas", "bayi": "Bayi 0-6 Bulan",
    "balita": "Balita & Anak Prasekolah", "remaja": "Usia Sekolah & Remaja",
    "dewasa": "Usia Dewasa", "lansia": "Lansia", "tbc": "Skrining TBC",
    "belum_ditentukan": "Belum Ditentukan",
}

# ---------------- Startup: seed ----------------
@app.on_event("startup")
async def startup():
    await db.users.create_index("username", unique=True)
    from seed_demo import seed_all
    await seed_all(db, hash_password)

# ==================== AUTH ====================
@api.post("/auth/login")
async def login(body: LoginIn, response: Response):
    user = await db.users.find_one({"username": body.username.lower().strip()})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(401, "Username atau kata sandi salah")
    if not user.get("aktif", True):
        raise HTTPException(403, "Akun dinonaktifkan")
    at = create_token(user, "access")
    rt = create_token(user, "refresh")
    for k, v, age in [("access_token", at, 43200), ("refresh_token", rt, 604800)]:
        response.set_cookie(k, v, httponly=True, secure=True, samesite="none", max_age=age, path="/")
    user.pop("_id", None); user.pop("password_hash", None)
    return {"user": user, "access_token": at}

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"ok": True}

@api.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return user

class ChangePasswordIn(BaseModel):
    current_password: str
    new_password: str

@api.post("/auth/change-password")
async def change_password(body: ChangePasswordIn, user=Depends(get_current_user)):
    u = await db.users.find_one({"id": user["id"]})
    if not u or not verify_password(body.current_password, u["password_hash"]):
        raise HTTPException(400, "Kata sandi saat ini salah")
    if len(body.new_password) < 6:
        raise HTTPException(400, "Kata sandi baru minimal 6 karakter")
    await db.users.update_one({"id": user["id"]}, {"$set": {"password_hash": hash_password(body.new_password)}})
    await audit(user, "update", "password", user["id"], "ganti kata sandi")
    return {"ok": True}

# ==================== MODE PIN ====================
class PinIn(BaseModel):
    pin: str

class ChangePinIn(BaseModel):
    current_pin: str
    new_pin: str

async def mode_pin_hash() -> str:
    s = await db.settings.find_one({"key": "mode_pin"})
    if s:
        return s["value"]
    h = hash_password(os.environ["MODE_PIN"])
    await db.settings.insert_one({"key": "mode_pin", "value": h})
    return h

@api.post("/mode/verify-pin")
async def verify_mode_pin(body: PinIn, user=Depends(require_admin)):
    if not verify_password(body.pin, await mode_pin_hash()):
        raise HTTPException(401, "PIN salah")
    return {"ok": True}

@api.post("/mode/change-pin")
async def change_mode_pin(body: ChangePinIn, user=Depends(require_admin)):
    if not verify_password(body.current_pin, await mode_pin_hash()):
        raise HTTPException(400, "PIN saat ini salah")
    if not (body.new_pin.isdigit() and 4 <= len(body.new_pin) <= 8):
        raise HTTPException(400, "PIN baru harus 4-8 digit angka")
    await db.settings.update_one({"key": "mode_pin"}, {"$set": {"value": hash_password(body.new_pin)}})
    await audit(user, "update", "mode_pin", "", "ganti PIN mode")
    return {"ok": True}

# ==================== MASTER DATA ====================
@api.get("/master/wilayah")
async def list_wilayah(user=Depends(get_current_user)):
    rows = await db.wilayah.find({"deleted": {"$ne": True}}, {"_id": 0}).to_list(500)
    return rows

@api.post("/master/wilayah")
async def add_wilayah(body: dict, user=Depends(require_admin)):
    doc = {"id": new_id(), "nama": body["nama"], "tipe": body.get("tipe", "kelurahan"),
           "parent": body.get("parent"), "deleted": False, "created_at": iso()}
    await db.wilayah.insert_one(doc); await audit(user, "create", "wilayah", doc["id"], body["nama"])
    doc.pop("_id", None); return doc

@api.get("/master/posyandu")
async def list_posyandu(user=Depends(get_current_user)):
    return await db.posyandu.find({"deleted": {"$ne": True}}, {"_id": 0}).to_list(500)

@api.get("/master/questions")
async def get_questions(group: Optional[str] = None, user=Depends(get_current_user)):
    q = {"deleted": {"$ne": True}}
    if group:
        q["group"] = group
    rows = await db.master_questions.find(q, {"_id": 0}).sort("urutan", 1).to_list(1000)
    return rows

@api.put("/master/questions/{kode}")
async def update_question(kode: str, body: dict, user=Depends(require_admin)):
    body.pop("_id", None); body.pop("kode", None); body.pop("group", None)
    body.pop("urutan", None); body.pop("custom", None); body.pop("deleted", None)
    if "section" in body and body["section"] not in QUESTION_SECTIONS:
        body.pop("section")
    if "jenis" in body and body["jenis"] not in QUESTION_JENIS:
        body.pop("jenis")
    if "opsi" in body and isinstance(body["opsi"], str):
        body["opsi"] = [s.strip() for s in body["opsi"].split(",") if s.strip()]
    if "problem_when" in body and isinstance(body["problem_when"], str):
        body["problem_when"] = [s.strip() for s in body["problem_when"].split(",") if s.strip()]
    if "satuan" in body:
        body["satuan"] = (str(body["satuan"]).strip() or None) if body["satuan"] else None
    if "priority" in body and body["priority"] not in ("kuning", "merah"):
        body["priority"] = None
    await db.master_questions.update_one({"kode": kode}, {"$set": body})
    await audit(user, "update", "master_pertanyaan", kode, json.dumps(body)[:200])
    return await db.master_questions.find_one({"kode": kode}, {"_id": 0})

QUESTION_SECTIONS = {"ceklis", "tanda_bahaya"}
QUESTION_JENIS = {"number", "single", "multi", "yesno", "pemeriksaan", "imunisasi", "danger"}

@api.post("/master/questions")
async def create_question(body: dict, user=Depends(require_admin)):
    text = str(body.get("text") or "").strip()
    group = str(body.get("group") or "").strip()
    if not text or not group:
        raise HTTPException(400, "Teks pertanyaan dan kelompok wajib diisi")
    if group not in KELOMPOK_LABEL or group == "belum_ditentukan":
        raise HTTPException(400, "Kelompok tidak valid")
    section = body.get("section") if body.get("section") in QUESTION_SECTIONS else "ceklis"
    jenis = body.get("jenis") if body.get("jenis") in QUESTION_JENIS else "yesno"
    kode = str(body.get("kode") or "").strip().upper()
    if kode:
        if await db.master_questions.find_one({"kode": kode}):
            raise HTTPException(400, "Kode sudah dipakai, gunakan kode lain")
    else:
        base = re.sub(r"[^A-Z0-9]+", "_", f"{group}_{text[:20]}".upper()).strip("_") or group.upper()
        kode = base; n = 1
        while await db.master_questions.find_one({"kode": kode}):
            n += 1; kode = f"{base}_{n}"
    last = await db.master_questions.find_one({"group": group}, sort=[("urutan", -1)])
    urutan = (last.get("urutan", 0) + 1) if last else 1
    opsi = body.get("opsi") or []
    if isinstance(opsi, str):
        opsi = [s.strip() for s in opsi.split(",") if s.strip()]
    problem_when = body.get("problem_when") or []
    if isinstance(problem_when, str):
        problem_when = [s.strip() for s in problem_when.split(",") if s.strip()]
    doc = {
        "kode": kode, "group": group, "section": section, "text": text,
        "definisi": str(body.get("definisi") or "").strip(),
        "jenis": jenis, "satuan": (str(body.get("satuan")).strip() if body.get("satuan") else None),
        "opsi": opsi, "wajib": bool(body.get("wajib", False)),
        "problem_when": problem_when, "priority": (body.get("priority") or None),
        "report_required": bool(body.get("report_required", False)),
        "urutan": urutan, "deleted": False, "custom": True,
    }
    await db.master_questions.insert_one(doc)
    await audit(user, "create", "master_pertanyaan", kode, text[:200])
    doc.pop("_id", None)
    return doc

@api.delete("/master/questions/{kode}")
async def delete_question(kode: str, user=Depends(require_admin)):
    r = await db.master_questions.update_one({"kode": kode}, {"$set": {"deleted": True}})
    if r.matched_count == 0:
        raise HTTPException(404, "Pertanyaan tidak ditemukan")
    await audit(user, "delete", "master_pertanyaan", kode, "")
    return {"ok": True}

@api.get("/master/groups")
async def groups(user=Depends(get_current_user)):
    return [{"code": k, "label": v} for k, v in KELOMPOK_LABEL.items() if k not in ("belum_ditentukan",)]

# ==================== KADER: KELUARGA ====================
def region_filter(user):
    if user["role"] == "admin":
        return {}
    # kader melihat keluarga di wilayahnya ATAU yang ditujukan langsung ke dirinya
    return {"$or": [{"kelurahan": {"$in": user.get("wilayah", [])}},
                    {"kader_id": user["id"]}]}

@api.get("/keluarga")
async def list_keluarga(search: str = "", kelurahan: str = "", page: int = 1, limit: int = 50, user=Depends(get_current_user)):
    q = {"deleted": {"$ne": True}}
    conds = []
    if user["role"] != "admin":
        conds.append(region_filter(user))
    if kelurahan:
        q["kelurahan"] = kelurahan
    if search:
        conds.append({"$or": [{"nama_kk": {"$regex": search, "$options": "i"}},
                    {"alamat": {"$regex": search, "$options": "i"}},
                    {"rt": search}, {"no_hp": {"$regex": search}}]})
    if conds:
        q["$and"] = conds
    total = await db.keluarga.count_documents(q)
    rows = await db.keluarga.find(q, {"_id": 0}).sort("nama_kk", 1).skip((page-1)*limit).limit(limit).to_list(limit)
    for r in rows:
        r["jumlah_anggota_aktual"] = await db.anggota.count_documents({"keluarga_id": r["id"], "deleted": {"$ne": True}})
        r["sudah_dikunjungi"] = await db.kunjungan.count_documents({"keluarga_id": r["id"], "status": {"$in": ["selesai", "terkirim"]}}) > 0
    return {"total": total, "items": rows, "page": page}

@api.post("/keluarga")
async def create_keluarga(body: KeluargaIn, user=Depends(get_current_user)):
    if user["role"] == "kader" and body.kelurahan not in user.get("wilayah", []):
        raise HTTPException(403, "Kelurahan di luar wilayah tugas Anda")
    doc = body.dict()
    doc.update({"id": new_id(), "deleted": False, "created_by": user["id"],
                "created_by_nama": user["nama"], "created_at": iso(),
                "puskesmas": "UPT Puskesmas Melati", "data_lengkap": bool(body.nama_kk)})
    await db.keluarga.insert_one(doc); await audit(user, "create", "keluarga", doc["id"], body.nama_kk)
    doc.pop("_id", None); return doc

@api.get("/keluarga/duplicate-check")
async def dup_check(nama: str = "", user=Depends(get_current_user)):
    q = {"deleted": {"$ne": True}, "nama_kk": {"$regex": nama, "$options": "i"}}
    q.update(region_filter(user))
    rows = await db.keluarga.find(q, {"_id": 0, "id": 1, "nama_kk": 1, "alamat": 1, "rt": 1, "rw": 1}).limit(5).to_list(5)
    return rows

@api.get("/keluarga/{kid}")
async def get_keluarga(kid: str, user=Depends(get_current_user)):
    k = await db.keluarga.find_one({"id": kid}, {"_id": 0})
    if not k:
        raise HTTPException(404, "Keluarga tidak ditemukan")
    anggota = await db.anggota.find({"keluarga_id": kid, "deleted": {"$ne": True}}, {"_id": 0}).to_list(100)
    for a in anggota:
        a["umur"] = hitung_umur(a.get("tanggal_lahir"))
        a["kelompok"] = tentukan_kelompok(a)
        a["kelompok_label"] = KELOMPOK_LABEL.get(a["kelompok"], a["kelompok"])
    k["anggota"] = anggota
    return k

@api.put("/keluarga/{kid}")
async def update_keluarga(kid: str, body: KeluargaIn, user=Depends(get_current_user)):
    doc = body.dict(); doc["data_lengkap"] = bool(body.nama_kk)
    await db.keluarga.update_one({"id": kid}, {"$set": doc})
    await audit(user, "update", "keluarga", kid, body.nama_kk)
    return await db.keluarga.find_one({"id": kid}, {"_id": 0})

@api.delete("/keluarga/{kid}")
async def delete_keluarga(kid: str, user=Depends(get_current_user)):
    k = await db.keluarga.find_one({"id": kid})
    if not k or k.get("deleted"):
        raise HTTPException(404, "Keluarga tidak ditemukan")
    if user["role"] == "kader" and k.get("kelurahan") not in user.get("wilayah", []):
        raise HTTPException(403, "Keluarga di luar wilayah tugas Anda")
    await db.keluarga.update_one({"id": kid}, {"$set": {"deleted": True, "deleted_at": iso(), "deleted_by": user["nama"]}})
    await db.anggota.update_many({"keluarga_id": kid}, {"$set": {"deleted": True}})
    await audit(user, "delete", "keluarga", kid, k.get("nama_kk", ""))
    return {"ok": True, "nama_kk": k.get("nama_kk", "")}

@api.post("/anggota")
async def add_anggota(body: AnggotaIn, user=Depends(get_current_user)):
    if body.nik and len(body.nik) != 16:
        pass  # allowed for draft, flagged below
    doc = body.dict()
    doc.update({"id": new_id(), "deleted": False, "created_at": iso(),
                "data_lengkap": bool(body.nik and len(body.nik) == 16 and body.tanggal_lahir)})
    await db.anggota.insert_one(doc); await audit(user, "create", "anggota", doc["id"], body.nama)
    doc.pop("_id", None)
    doc["umur"] = hitung_umur(doc.get("tanggal_lahir"))
    doc["kelompok"] = tentukan_kelompok(doc)
    doc["kelompok_label"] = KELOMPOK_LABEL.get(doc["kelompok"], doc["kelompok"])
    return doc

@api.put("/anggota/{aid}")
async def update_anggota(aid: str, body: AnggotaIn, user=Depends(get_current_user)):
    doc = body.dict()
    doc["data_lengkap"] = bool(body.nik and len(body.nik) == 16 and body.tanggal_lahir)
    await db.anggota.update_one({"id": aid}, {"$set": doc})
    await audit(user, "update", "anggota", aid, body.nama)
    out = await db.anggota.find_one({"id": aid}, {"_id": 0})
    out["umur"] = hitung_umur(out.get("tanggal_lahir")); out["kelompok"] = tentukan_kelompok(out)
    out["kelompok_label"] = KELOMPOK_LABEL.get(out["kelompok"], out["kelompok"])
    return out

# ==================== RULE ENGINE & VISIT ====================
async def evaluate_answers(anggota, answers):
    """Return list of temuan (findings) from answers using master_questions rules."""
    temuan = []
    codes = list(answers.keys())
    qs = await db.master_questions.find({"kode": {"$in": codes}}, {"_id": 0}).to_list(500)
    qmap = {q["kode"]: q for q in qs}
    for kode, val in answers.items():
        q = qmap.get(kode)
        if not q or not q.get("problem_when"):
            continue
        hit = False
        if isinstance(val, list):
            hit = any(v in q["problem_when"] for v in val)
        else:
            hit = val in q["problem_when"]
        if hit:
            temuan.append({
                "kode": kode, "masalah": q["text"], "definisi": q.get("definisi", ""),
                "priority": q.get("priority") or "kuning", "group": q["group"],
                "report_required": q.get("report_required", False),
            })
    return temuan

@api.post("/kunjungan")
async def create_kunjungan(body: KunjunganIn, user=Depends(get_current_user)):
    kel = await db.keluarga.find_one({"id": body.keluarga_id}, {"_id": 0})
    if not kel:
        raise HTTPException(404, "Keluarga tidak ditemukan")
    vid = new_id()
    all_temuan = []
    per_anggota = []
    for aid in body.anggota_ids:
        a = await db.anggota.find_one({"id": aid}, {"_id": 0})
        if not a:
            continue
        ans = body.answers.get(aid, {})
        temuan = await evaluate_answers(a, ans)
        for t in temuan:
            t["anggota_id"] = aid; t["anggota_nama"] = a["nama"]
        all_temuan.extend(temuan)
        per_anggota.append({"anggota_id": aid, "nama": a["nama"],
                            "kelompok": tentukan_kelompok(a), "answers": ans,
                            "edukasi": body.edukasi.get(aid, []), "temuan": temuan})
    has_red = any(t["priority"] == "merah" for t in all_temuan)
    if has_red and body.status == "selesai" and not body.reminder_confirmed:
        raise HTTPException(400, "Tanda bahaya terdeteksi. Konfirmasi bahwa Anda sudah mengingatkan sasaran & akan melaporkan ke petugas.")
    overall = "merah" if has_red else ("kuning" if all_temuan else "hijau")
    doc = {
        "id": vid, "keluarga_id": body.keluarga_id, "keluarga_nama": kel["nama_kk"],
        "kelurahan": kel["kelurahan"], "rt": kel.get("rt", ""), "rw": kel.get("rw", ""),
        "posyandu": kel.get("posyandu", ""), "anggota_ids": body.anggota_ids,
        "kader_id": user["id"], "kader_nama": user["nama"],
        "tanggal": body.tanggal or iso(), "catatan": body.catatan, "gps": body.gps,
        "per_anggota": per_anggota, "status": "terkirim" if body.status == "selesai" else "draf",
        "sync_status": "tersinkron", "prioritas": overall,
        "created_at": iso(), "reminder_confirmed": body.reminder_confirmed,
    }
    await db.kunjungan.insert_one(doc)
    await audit(user, "create", "kunjungan", vid, f"{kel['nama_kk']} - {overall}")
    # create follow-up cases + notifications for problems when submitted
    if doc["status"] == "terkirim":
        await create_cases(doc, all_temuan, user)
    doc.pop("_id", None)
    return doc

async def create_cases(visit, temuan, user):
    for t in temuan:
        cid = new_id()
        target = 24 if t["priority"] == "merah" else 72
        case = {
            "id": cid, "kunjungan_id": visit["id"], "keluarga_id": visit["keluarga_id"],
            "keluarga_nama": visit["keluarga_nama"], "anggota_id": t.get("anggota_id"),
            "sasaran_nama": t.get("anggota_nama", ""), "group": t["group"],
            "masalah": t["masalah"], "definisi": t["definisi"], "priority": t["priority"],
            "kelurahan": visit["kelurahan"], "rt": visit["rt"], "rw": visit["rw"],
            "posyandu": visit["posyandu"], "kader_nama": visit["kader_nama"],
            "waktu_lapor": iso(), "status": "baru", "penanggung_jawab": None,
            "target_selesai": iso(now_utc() + timedelta(hours=target)),
            "riwayat": [{"waktu": iso(), "aksi": "Kasus dibuat dari kunjungan", "oleh": visit["kader_nama"], "status": "baru"}],
            "created_at": iso(),
        }
        await db.kasus.insert_one(case)
        await db.notifikasi.insert_one({
            "id": new_id(), "tipe": "kasus_baru" if t["priority"] != "merah" else "prioritas_merah",
            "priority": t["priority"], "judul": f"Laporan {t['priority'].upper()}: {t['masalah']}",
            "pesan": f"Dari kader {visit['kader_nama']} di {visit['kelurahan']} RT {visit['rt']}",
            "kasus_id": cid, "dibaca": False, "untuk_role": "admin", "waktu": iso()})

@api.get("/kunjungan")
async def list_kunjungan(mine: bool = False, keluarga_id: str = "", status: str = "", user=Depends(get_current_user)):
    q = {}
    if user["role"] == "kader" or mine:
        q["kader_id"] = user["id"]
    if keluarga_id:
        q["keluarga_id"] = keluarga_id
    if status:
        q["status"] = status
    rows = await db.kunjungan.find(q, {"_id": 0, "per_anggota": 0}).sort("created_at", -1).limit(200).to_list(200)
    return rows

@api.get("/kunjungan/{vid}")
async def get_kunjungan(vid: str, user=Depends(get_current_user)):
    v = await db.kunjungan.find_one({"id": vid}, {"_id": 0})
    if not v:
        raise HTTPException(404, "Kunjungan tidak ditemukan")
    return v

# ==================== KADER: TUGAS & STATUS LAPORAN ====================
@api.get("/kader/beranda")
async def kader_beranda(user=Depends(get_current_user)):
    wil = user.get("wilayah", [])
    total_kel = await db.keluarga.count_documents({"kelurahan": {"$in": wil}, "deleted": {"$ne": True}})
    dikunjungi_ids = await db.kunjungan.distinct("keluarga_id", {"kader_id": user["id"], "status": "terkirim"})
    masalah = await db.kasus.count_documents({"kader_nama": user["nama"]})
    ditindak = await db.kasus.count_documents({"kader_nama": user["nama"], "status": {"$in": ["selesai", "dirujuk", "sudah_dikunjungi"]}})
    draf = await db.kunjungan.count_documents({"kader_id": user["id"], "status": "draf"})
    return {
        "nama": user["nama"], "posyandu": user.get("posyandu", ""), "wilayah": wil,
        "target": user.get("target_keluarga", total_kel), "total_keluarga": total_kel,
        "sudah_dikunjungi": len(dikunjungi_ids), "belum_dikunjungi": max(total_kel - len(dikunjungi_ids), 0),
        "masalah_dilaporkan": masalah, "sudah_ditindaklanjuti": ditindak, "draf_belum_terkirim": draf,
    }

@api.get("/kader/laporan-status")
async def laporan_status(user=Depends(get_current_user)):
    cases = await db.kasus.find({"kader_nama": user["nama"]}, {"_id": 0}).sort("created_at", -1).limit(100).to_list(100)
    return cases

# ==================== ADMIN: DASHBOARD ====================
def dt_range(start, end):
    q = {}
    if start:
        q["$gte"] = start
    if end:
        q["$lte"] = end + "T23:59:59"
    return q

@api.get("/dashboard/kpi")
async def dashboard_kpi(start: str = "", end: str = "", kelurahan: str = "", user=Depends(require_admin)):
    kq = {"deleted": {"$ne": True}}
    vq = {"status": "terkirim"}
    cq = {}
    if kelurahan:
        kq["kelurahan"] = kelurahan; vq["kelurahan"] = kelurahan; cq["kelurahan"] = kelurahan
    if start or end:
        vq["created_at"] = dt_range(start, end); cq["waktu_lapor"] = dt_range(start, end)
    total_kel = await db.keluarga.count_documents(kq)
    visited_ids = await db.kunjungan.distinct("keluarga_id", vq)
    total_visits = await db.kunjungan.count_documents(vq)
    sasaran = 0
    async for v in db.kunjungan.find(vq, {"anggota_ids": 1}):
        sasaran += len(v.get("anggota_ids", []))
    merah = await db.kasus.count_documents({**cq, "priority": "merah"})
    kuning = await db.kasus.count_documents({**cq, "priority": "kuning"})
    belum = await db.kasus.count_documents({**cq, "status": {"$in": ["baru", "sudah_dibaca"]}})
    proses = await db.kasus.count_documents({**cq, "status": {"$in": ["ditugaskan", "dihubungi", "dijadwalkan", "sudah_dikunjungi", "dirujuk", "menunggu_hasil"]}})
    selesai = await db.kasus.count_documents({**cq, "status": "selesai"})
    # median response time (jam) for resolved cases
    resp_times = []
    async for c in db.kasus.find({**cq, "status": "selesai"}, {"waktu_lapor": 1, "riwayat": 1}):
        for r in c.get("riwayat", []):
            if r.get("status") == "selesai":
                try:
                    t0 = datetime.fromisoformat(c["waktu_lapor"]); t1 = datetime.fromisoformat(r["waktu"])
                    resp_times.append((t1 - t0).total_seconds() / 3600)
                except Exception:
                    pass
    resp_times.sort()
    median = round(resp_times[len(resp_times)//2], 1) if resp_times else 0
    data_belum_lengkap = await db.anggota.count_documents({"data_lengkap": False, "deleted": {"$ne": True}})
    return {
        "keluarga_terdaftar": total_kel, "keluarga_dikunjungi": len(visited_ids),
        "cakupan": round(len(visited_ids)/total_kel*100, 1) if total_kel else 0,
        "total_kunjungan": total_visits, "sasaran_dikunjungi": sasaran,
        "kasus_merah": merah, "kasus_kuning": kuning, "belum_ditindaklanjuti": belum,
        "sedang_ditindaklanjuti": proses, "selesai": selesai,
        "median_respons_jam": median, "data_belum_lengkap": data_belum_lengkap,
    }

@api.get("/dashboard/charts")
async def dashboard_charts(start: str = "", end: str = "", kelurahan: str = "", user=Depends(require_admin)):
    vq = {"status": "terkirim"}
    cq = {}
    if kelurahan:
        vq["kelurahan"] = kelurahan; cq["kelurahan"] = kelurahan
    if start or end:
        vq["created_at"] = dt_range(start, end)
    # coverage per kelurahan
    per_kel = defaultdict(lambda: {"kunjungan": 0})
    async for v in db.kunjungan.find(vq, {"kelurahan": 1}):
        per_kel[v.get("kelurahan", "-")]["kunjungan"] += 1
    cov = [{"kelurahan": k, "kunjungan": val["kunjungan"]} for k, val in per_kel.items()]
    # distribusi kelompok sasaran
    kelompok = Counter()
    async for v in db.kunjungan.find(vq, {"per_anggota": 1}):
        for pa in v.get("per_anggota", []):
            kelompok[pa.get("kelompok", "-")] += 1
    dist = [{"kelompok": KELOMPOK_LABEL.get(k, k), "jumlah": c} for k, c in kelompok.items()]
    # top masalah
    masalah = Counter()
    async for c in db.kasus.find(cq, {"masalah": 1}):
        masalah[c["masalah"]] += 1
    top = [{"masalah": m, "jumlah": c} for m, c in masalah.most_common(10)]
    # tren bulanan
    tren = defaultdict(int)
    async for v in db.kunjungan.find(vq, {"created_at": 1}):
        tren[str(v["created_at"])[:7]] += 1
    trend = [{"bulan": k, "jumlah": tren[k]} for k in sorted(tren)]
    # status tindak lanjut
    stat = Counter()
    async for c in db.kasus.find(cq, {"status": 1}):
        stat[c["status"]] += 1
    status_tl = [{"status": k, "jumlah": v} for k, v in stat.items()]
    return {"cakupan_kelurahan": cov, "distribusi_kelompok": dist, "top_masalah": top,
            "tren_bulanan": trend, "status_tindak_lanjut": status_tl}

# ==================== ADMIN: KASUS ====================
CASE_STATUS = ["baru", "sudah_dibaca", "ditugaskan", "dihubungi", "dijadwalkan",
               "sudah_dikunjungi", "dirujuk", "menunggu_hasil", "selesai", "tidak_dapat_ditindaklanjuti"]

@api.get("/kasus")
async def list_kasus(priority: str = "", status: str = "", kelurahan: str = "", group: str = "",
                     sort: str = "priority", user=Depends(require_admin)):
    q = {}
    for f, v in [("priority", priority), ("status", status), ("kelurahan", kelurahan), ("group", group)]:
        if v:
            q[f] = v
    rows = await db.kasus.find(q, {"_id": 0}).to_list(500)
    prio_rank = {"merah": 0, "kuning": 1, "hijau": 2}
    if sort == "priority":
        rows.sort(key=lambda r: (prio_rank.get(r["priority"], 9), r["waktu_lapor"]))
    else:
        rows.sort(key=lambda r: r["waktu_lapor"])
    for r in rows:
        r["nik_masked"] = "****"
    return rows

@api.get("/kasus/{cid}")
async def get_kasus(cid: str, user=Depends(require_admin)):
    c = await db.kasus.find_one({"id": cid}, {"_id": 0})
    if not c:
        raise HTTPException(404, "Kasus tidak ditemukan")
    if c["status"] == "baru":
        await db.kasus.update_one({"id": cid}, {"$set": {"status": "sudah_dibaca"},
            "$push": {"riwayat": {"waktu": iso(), "aksi": "Kasus dibaca petugas", "oleh": user["nama"], "status": "sudah_dibaca"}}})
        c = await db.kasus.find_one({"id": cid}, {"_id": 0})
    return c

@api.put("/kasus/{cid}")
async def update_kasus(cid: str, body: CaseUpdateIn, user=Depends(require_admin)):
    c = await db.kasus.find_one({"id": cid})
    if not c:
        raise HTTPException(404, "Kasus tidak ditemukan")
    if body.status == "tidak_dapat_ditindaklanjuti" and not body.alasan:
        raise HTTPException(400, "Alasan wajib diisi untuk status 'Tidak dapat ditindaklanjuti'")
    upd = {k: v for k, v in body.dict().items() if v is not None and k != "alasan"}
    hist = {"waktu": iso(), "oleh": user["nama"], "aksi": "Pembaruan kasus", "status": body.status or c["status"]}
    if body.hasil:
        hist["hasil"] = body.hasil
    if body.alasan:
        upd["alasan"] = body.alasan; hist["alasan"] = body.alasan
    await db.kasus.update_one({"id": cid}, {"$set": upd, "$push": {"riwayat": hist}})
    await audit(user, "update", "kasus", cid, body.status or "")
    return await db.kasus.find_one({"id": cid}, {"_id": 0})

# ==================== ADMIN: TINDAK LANJUT PUSTU ====================
@api.get("/admin/tindak-lanjut/sumber")
async def tl_sumber(posyandu: str = "", user=Depends(require_admin)):
    """Daftar kasus temuan kunjungan sebagai sumber tindak lanjut Pustu (auto-isi identitas & masalah)."""
    q = {}
    if posyandu:
        q["posyandu"] = posyandu
    kasus = await db.kasus.find(q, {"_id": 0}).to_list(1000)
    prio_rank = {"merah": 0, "kuning": 1, "hijau": 2}
    kasus.sort(key=lambda r: (prio_rank.get(r.get("priority"), 9), r.get("waktu_lapor", "")))
    done_ids = set(await db.tindak_lanjut_pustu.distinct("kasus_id"))
    out = []
    for k in kasus:
        a = await db.anggota.find_one({"id": k.get("anggota_id")}, {"_id": 0}) if k.get("anggota_id") else None
        kel = await db.keluarga.find_one({"id": k.get("keluarga_id")}, {"_id": 0}) if k.get("keluarga_id") else None
        a = a or {}; kel = kel or {}
        alamat = kel.get("alamat", "") or ""
        if kel.get("rt") or kel.get("rw"):
            alamat = f"{alamat} RT {kel.get('rt','-')}/{kel.get('rw','-')}".strip()
        out.append({
            "kasus_id": k["id"], "posyandu": k.get("posyandu", ""), "kelurahan": k.get("kelurahan", ""),
            "priority": k.get("priority", ""), "status": k.get("status", ""),
            "nama": k.get("sasaran_nama", ""), "masalah": k.get("masalah", ""),
            "nik": a.get("nik", "") or "", "tanggal_lahir": (a.get("tanggal_lahir") or "")[:10],
            "no_telp": a.get("no_hp", "") or kel.get("no_hp", "") or "",
            "alamat": alamat, "sudah_ada_tl": k["id"] in done_ids,
        })
    return out

@api.get("/admin/tindak-lanjut")
async def tl_list(posyandu: str = "", user=Depends(require_admin)):
    q = {}
    if posyandu:
        q["posyandu"] = posyandu
    return await db.tindak_lanjut_pustu.find(q, {"_id": 0}).sort("created_at", -1).limit(1000).to_list(1000)

@api.post("/admin/tindak-lanjut")
async def tl_create(body: dict, user=Depends(require_admin)):
    nama = str(body.get("nama") or "").strip()
    posyandu = str(body.get("posyandu") or "").strip()
    tindak_lanjut = str(body.get("tindak_lanjut") or "").strip()
    if not nama or not posyandu or not tindak_lanjut:
        raise HTTPException(400, "Nama, posyandu, dan tindak lanjut wajib diisi")
    doc = {
        "id": new_id(), "kasus_id": body.get("kasus_id") or None,
        "posyandu": posyandu, "nama": nama,
        "nik": str(body.get("nik") or "").strip(),
        "tanggal_lahir": str(body.get("tanggal_lahir") or "").strip(),
        "alamat": str(body.get("alamat") or "").strip(),
        "no_telp": str(body.get("no_telp") or "").strip(),
        "masalah": str(body.get("masalah") or "").strip(),
        "tindak_lanjut": tindak_lanjut,
        "waktu": str(body.get("waktu") or "").strip() or iso(),
        "petugas": str(body.get("petugas") or "").strip(),
        "kelurahan": str(body.get("kelurahan") or "").strip(),
        "created_by": user["nama"], "created_at": iso(),
    }
    await db.tindak_lanjut_pustu.insert_one(doc)
    await audit(user, "create", "tindak_lanjut_pustu", doc["id"], f"{nama} - {posyandu}")
    doc.pop("_id", None)
    return doc

@api.put("/admin/tindak-lanjut/{tid}")
async def tl_update(tid: str, body: dict, user=Depends(require_admin)):
    upd = {}
    for f in ["posyandu", "nama", "nik", "tanggal_lahir", "alamat", "no_telp", "masalah", "tindak_lanjut", "waktu", "petugas", "kelurahan"]:
        if f in body:
            upd[f] = str(body.get(f) or "").strip()
    if not upd:
        raise HTTPException(400, "Tidak ada perubahan")
    r = await db.tindak_lanjut_pustu.update_one({"id": tid}, {"$set": upd})
    if r.matched_count == 0:
        raise HTTPException(404, "Data tidak ditemukan")
    await audit(user, "update", "tindak_lanjut_pustu", tid, "")
    return await db.tindak_lanjut_pustu.find_one({"id": tid}, {"_id": 0})

@api.delete("/admin/tindak-lanjut/{tid}")
async def tl_delete(tid: str, user=Depends(require_admin)):
    r = await db.tindak_lanjut_pustu.delete_one({"id": tid})
    if r.deleted_count == 0:
        raise HTTPException(404, "Data tidak ditemukan")
    await audit(user, "delete", "tindak_lanjut_pustu", tid, "")
    return {"ok": True}


# ==================== NOTIFIKASI ====================
@api.get("/notifikasi")
async def list_notif(user=Depends(require_admin)):
    rows = await db.notifikasi.find({"untuk_role": "admin"}, {"_id": 0}).sort("waktu", -1).limit(50).to_list(50)
    unread = await db.notifikasi.count_documents({"untuk_role": "admin", "dibaca": False})
    return {"items": rows, "unread": unread}

@api.post("/notifikasi/{nid}/read")
async def read_notif(nid: str, user=Depends(require_admin)):
    await db.notifikasi.update_one({"id": nid}, {"$set": {"dibaca": True}})
    return {"ok": True}

# ==================== REKAP ====================
@api.get("/rekap")
async def rekap(periode: str = "bulan", user=Depends(require_admin)):
    vq = {"status": "terkirim"}
    kelompok = Counter(); tb = 0; edukasi = 0; dilaporkan = 0
    unique_sasaran = set()
    async for v in db.kunjungan.find(vq, {"per_anggota": 1, "prioritas": 1}):
        for pa in v.get("per_anggota", []):
            kelompok[pa.get("kelompok")] += 1
            if pa.get("edukasi"):
                edukasi += 1
            if pa.get("temuan"):
                unique_sasaran.add(pa.get("anggota_id"))
                dilaporkan += 1
                if any(t["priority"] == "merah" for t in pa["temuan"]):
                    tb += 1
    kel_visited = len(await db.kunjungan.distinct("keluarga_id", vq))
    selesai = await db.kasus.count_documents({"status": "selesai"})
    belum = await db.kasus.count_documents({"status": {"$ne": "selesai"}})
    return {
        "keluarga_dikunjungi": kel_visited,
        "per_kelompok": [{"kelompok": KELOMPOK_LABEL.get(k, k), "jumlah": c} for k, c in kelompok.items()],
        "tanda_bahaya": tb, "edukasi": edukasi, "sasaran_bermasalah": len(unique_sasaran),
        "dilaporkan": dilaporkan, "kasus_selesai": selesai, "kasus_belum_selesai": belum,
    }

# ==================== EXPORT ====================
REPORT_TITLES = {
    "kasus": "DAFTAR TINDAK LANJUT & SASARAN BERMASALAH",
    "kunjungan": "REKAP KUNJUNGAN RUMAH KADER",
    "keluarga": "DAFTAR KELUARGA TERDAFTAR",
    "tindak_lanjut": "REKAP TINDAK LANJUT PUSTU",
}
EXPORT_STATUS_LABEL = {
    "baru": "Baru", "sudah_dibaca": "Sudah Dibaca", "ditugaskan": "Ditugaskan", "dihubungi": "Dihubungi",
    "dijadwalkan": "Dijadwalkan", "sudah_dikunjungi": "Sudah Dikunjungi", "dirujuk": "Dirujuk",
    "menunggu_hasil": "Menunggu Hasil", "selesai": "Selesai", "tidak_dapat_ditindaklanjuti": "Tidak Dapat Ditindaklanjuti",
    "terkirim": "Terkirim", "draf": "Draf",
}
EXPORT_PRIORITY_LABEL = {"merah": "Merah (Tanda Bahaya)", "kuning": "Kuning (Perlu Perhatian)", "hijau": "Hijau (Normal)"}


def _fmt_tgl(v, with_time=False):
    if not v:
        return "-"
    s = str(v)
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
        return d.strftime("%d-%m-%Y %H:%M") if with_time else d.strftime("%d-%m-%Y")
    except Exception:
        return s[:10] or "-"


def _yn(v):
    return "Ya" if v is True else ("Tidak" if v is False else "-")


def _rtrw(r):
    rt, rw = r.get("rt") or "-", r.get("rw") or "-"
    return f"{rt}/{rw}"


def _report_columns(jenis):
    """Return list of (label, extractor) with curated, human-readable columns."""
    if jenis == "kasus":
        return [
            ("Nama Sasaran", lambda r: r.get("sasaran_nama") or "-"),
            ("Kelompok", lambda r: KELOMPOK_LABEL.get(r.get("group"), r.get("group") or "-")),
            ("Masalah Ditemukan", lambda r: r.get("masalah") or "-"),
            ("Prioritas", lambda r: EXPORT_PRIORITY_LABEL.get(r.get("priority"), r.get("priority") or "-")),
            ("Kelurahan", lambda r: r.get("kelurahan") or "-"),
            ("RT/RW", _rtrw),
            ("Posyandu", lambda r: r.get("posyandu") or "-"),
            ("Kader Pelapor", lambda r: r.get("kader_nama") or "-"),
            ("Waktu Lapor", lambda r: _fmt_tgl(r.get("waktu_lapor"), True)),
            ("Status", lambda r: EXPORT_STATUS_LABEL.get(r.get("status"), r.get("status") or "-")),
            ("Target Selesai", lambda r: _fmt_tgl(r.get("target_selesai"), True)),
        ]
    if jenis == "kunjungan":
        return [
            ("Tanggal", lambda r: _fmt_tgl(r.get("tanggal") or r.get("created_at"))),
            ("Kepala Keluarga", lambda r: r.get("keluarga_nama") or "-"),
            ("Kelurahan", lambda r: r.get("kelurahan") or "-"),
            ("RT/RW", _rtrw),
            ("Posyandu", lambda r: r.get("posyandu") or "-"),
            ("Kader", lambda r: r.get("kader_nama") or "-"),
            ("Anggota Dikunjungi", lambda r: str(len(r.get("anggota_ids") or []))),
            ("Prioritas", lambda r: EXPORT_PRIORITY_LABEL.get(r.get("prioritas"), r.get("prioritas") or "-")),
            ("Status", lambda r: EXPORT_STATUS_LABEL.get(r.get("status"), r.get("status") or "-")),
        ]
    if jenis == "keluarga":
        return [
            ("Nama Kepala Keluarga", lambda r: r.get("nama_kk") or "-"),
            ("Alamat", lambda r: r.get("alamat") or "-"),
            ("RT/RW", _rtrw),
            ("Kelurahan", lambda r: r.get("kelurahan") or "-"),
            ("Posyandu", lambda r: r.get("posyandu") or "-"),
            ("No. HP", lambda r: r.get("no_hp") or "-"),
            ("Jml Anggota", lambda r: str(r.get("jumlah_anggota") or "-")),
            ("JKN", lambda r: _yn(r.get("punya_jkn"))),
            ("Air Bersih", lambda r: _yn(r.get("air_bersih"))),
            ("Jamban", lambda r: _yn(r.get("jamban"))),
            ("Ventilasi", lambda r: _yn(r.get("ventilasi"))),
        ]
    if jenis == "tindak_lanjut":
        return [
            ("Posyandu", lambda r: r.get("posyandu") or "-"),
            ("Nama", lambda r: r.get("nama") or "-"),
            ("NIK", lambda r: r.get("nik") or "-"),
            ("Tanggal Lahir", lambda r: _fmt_tgl(r.get("tanggal_lahir"))),
            ("Alamat", lambda r: r.get("alamat") or "-"),
            ("No. Telp", lambda r: r.get("no_telp") or "-"),
            ("Masalah Ditemukan", lambda r: r.get("masalah") or "-"),
            ("Tindak Lanjut", lambda r: r.get("tindak_lanjut") or "-"),
            ("Petugas", lambda r: r.get("petugas") or "-"),
            ("Waktu", lambda r: _fmt_tgl(r.get("waktu"))),
        ]
    return None


async def _report_rows(jenis, start, end, kelurahan):
    if jenis == "kasus":
        q = {}
        if kelurahan:
            q["kelurahan"] = kelurahan
        if start or end:
            q["waktu_lapor"] = dt_range(start, end)
        return await db.kasus.find(q, {"_id": 0, "riwayat": 0}).sort("waktu_lapor", -1).to_list(5000)
    if jenis == "keluarga":
        q = {"deleted": {"$ne": True}}
        if kelurahan:
            q["kelurahan"] = kelurahan
        if start or end:
            q["created_at"] = dt_range(start, end)
        return await db.keluarga.find(q, {"_id": 0}).sort("nama_kk", 1).to_list(5000)
    if jenis == "kunjungan":
        q = {}
        if kelurahan:
            q["kelurahan"] = kelurahan
        if start or end:
            q["created_at"] = dt_range(start, end)
        return await db.kunjungan.find(q, {"_id": 0, "per_anggota": 0}).sort("created_at", -1).to_list(5000)
    if jenis == "tindak_lanjut":
        q = {}
        if kelurahan:
            q["kelurahan"] = kelurahan
        if start or end:
            q["created_at"] = dt_range(start, end)
        return await db.tindak_lanjut_pustu.find(q, {"_id": 0}).sort("created_at", -1).to_list(5000)
    return None


def _periode_str(start, end, kelurahan):
    if start or end:
        per = f"Periode: {_fmt_tgl(start) if start else 'Awal'} s.d. {_fmt_tgl(end) if end else 'Sekarang'}"
    else:
        per = "Periode: Semua Data"
    if kelurahan:
        per += f"  |  Kelurahan: {kelurahan}"
    return per


@api.get("/export/{jenis}")
async def export(jenis: str, fmt: str = "csv", start: str = "", end: str = "", kelurahan: str = "", user=Depends(require_admin)):
    cols = _report_columns(jenis)
    if cols is None:
        raise HTTPException(400, "Jenis laporan tidak dikenal")
    rows = await _report_rows(jenis, start, end, kelurahan)
    ak = await db.akreditasi.find_one({}, {"_id": 0, "puskesmas": 1})
    puskesmas = (ak or {}).get("puskesmas") or "UPT Puskesmas Melati"
    title = REPORT_TITLES[jenis]
    periode = _periode_str(start, end, kelurahan)
    dicetak = f"Dicetak: {_fmt_tgl(iso(), True)} oleh {user['nama']}"
    headers = ["No"] + [c[0] for c in cols]
    matrix = []
    for i, r in enumerate(rows, 1):
        matrix.append([str(i)] + [c[1](r) for c in cols])

    await audit(user, "export", jenis, "", f"{fmt} {start}-{end} {kelurahan}".strip())
    fname = f"laporan_{jenis}_{now_utc().strftime('%Y%m%d')}"

    if fmt == "csv":
        import csv
        buf = io.StringIO()
        buf.write(f"{title}\r\n{puskesmas}\r\n{periode}\r\n{dicetak}\r\n\r\n")
        w = csv.writer(buf)
        w.writerow(headers)
        for row in matrix:
            w.writerow(row)
        w.writerow([])
        w.writerow(["", f"Total data: {len(matrix)}"])
        return StreamingResponse(io.BytesIO(buf.getvalue().encode("utf-8-sig")), media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={fname}.csv"})

    if fmt == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = jenis[:28]
        ncol = len(headers)
        last_col = get_column_letter(ncol)
        thin = Side(style="thin", color="B0B0B0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        # kop
        meta = [(title, 14, True), (puskesmas, 11, True), (periode, 9, False), (dicetak, 9, False)]
        for ri, (txt, sz, bold) in enumerate(meta, 1):
            ws.merge_cells(f"A{ri}:{last_col}{ri}")
            c = ws.cell(row=ri, column=1, value=txt)
            c.font = Font(size=sz, bold=bold, color="0F766E" if ri == 1 else "334155")
            c.alignment = Alignment(horizontal="center" if ri <= 2 else "left", vertical="center")
        head_row = len(meta) + 2
        # header
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=head_row, column=ci, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F766E")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        # data
        for ri, row in enumerate(matrix, head_row + 1):
            for ci, val in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.alignment = Alignment(vertical="top", wrap_text=True, horizontal="center" if ci == 1 else "left")
                c.border = border
        # column widths
        widths = [len(h) for h in headers]
        for row in matrix:
            for ci, val in enumerate(row):
                widths[ci] = max(widths[ci], min(len(str(val)), 45))
        for ci, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = max(6, min(wd + 3, 48))
        ws.freeze_panes = f"A{head_row + 1}"
        ws.cell(row=head_row + len(matrix) + 2, column=1, value=f"Total data: {len(matrix)}").font = Font(bold=True, italic=True)
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={fname}.xlsx"})

    if fmt == "pdf":
        from fpdf import FPDF

        def latin(s):
            return str(s).encode("latin-1", "replace").decode("latin-1")

        weights = {
            "No": 0.5, "Masalah Ditemukan": 3.2, "Tindak Lanjut": 3.2, "Alamat": 2.6,
            "Nama Sasaran": 2.2, "Kepala Keluarga": 2.2, "Nama Kepala Keluarga": 2.4, "Nama": 2.2,
            "Kelompok": 1.6, "Prioritas": 1.8, "Status": 1.8, "Kader Pelapor": 1.8, "Kader": 1.8,
            "Waktu Lapor": 1.6, "Target Selesai": 1.6, "Tanggal": 1.2, "Waktu": 1.2, "Tanggal Lahir": 1.4,
            "NIK": 1.8, "No. Telp": 1.4, "No. HP": 1.4, "Petugas": 1.6,
        }
        pdf = FPDF(orientation="L", unit="mm", format="A4")
        pdf.set_auto_page_break(False)
        pdf.set_margins(10, 10, 10)
        epw = pdf.epw
        wsum = sum(weights.get(h, 1.4) for h in headers)
        colw = [epw * weights.get(h, 1.4) / wsum for h in headers]
        line_h = 4.6
        bottom_limit = pdf.h - 18

        def header_kop():
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(epw, 6, latin(title), align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(epw, 5, latin(puskesmas), align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 8)
            pdf.cell(epw, 4.5, latin(periode), align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.cell(epw, 4.5, latin(dicetak), align="C", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            draw_head()

        def draw_head():
            pdf.set_font("Helvetica", "B", 7.5)
            pdf.set_fill_color(15, 118, 110)
            pdf.set_text_color(255, 255, 255)
            x, y = pdf.get_x(), pdf.get_y()
            for i, h in enumerate(headers):
                pdf.multi_cell(colw[i], 6, latin(h), border=1, align="C", fill=True,
                               max_line_height=3, new_x="RIGHT", new_y="TOP")
            pdf.set_xy(x, y + 6)
            pdf.set_text_color(30, 41, 59)

        def wrap(txt, w):
            pdf.set_font("Helvetica", "", 7)
            words = str(txt).split()
            lines, cur = [], ""
            for wd in words:
                t = (cur + " " + wd).strip()
                if pdf.get_string_width(latin(t)) <= w - 2:
                    cur = t
                else:
                    if cur:
                        lines.append(cur)
                    cur = wd
            if cur:
                lines.append(cur)
            return lines or [""]

        header_kop()
        pdf.set_font("Helvetica", "", 7)
        fill = False
        for row in matrix:
            cell_lines = [wrap(row[i], colw[i]) for i in range(len(headers))]
            nlines = max(len(cl) for cl in cell_lines)
            rh = nlines * line_h + 1.5
            if pdf.get_y() + rh > bottom_limit:
                pdf.set_font("Helvetica", "", 7)
                header_kop()
                pdf.set_font("Helvetica", "", 7)
            x, y = pdf.get_x(), pdf.get_y()
            pdf.set_fill_color(244, 247, 246)
            for i in range(len(headers)):
                pdf.rect(x, y, colw[i], rh, style="DF" if fill else "D")
                ty = y + 1
                for ln in cell_lines[i]:
                    pdf.set_xy(x + 1, ty)
                    pdf.cell(colw[i] - 2, line_h, latin(ln), align="C" if i == 0 else "L")
                    ty += line_h
                x += colw[i]
            pdf.set_xy(pdf.l_margin, y + rh)
            fill = not fill

        pdf.ln(2)
        pdf.set_font("Helvetica", "B", 8)
        if pdf.get_y() + 40 > pdf.h - 10:
            pdf.add_page()
        pdf.cell(epw, 5, latin(f"Total data: {len(matrix)}"), new_x="LMARGIN", new_y="NEXT")
        # tanda tangan
        pdf.ln(4)
        pdf.set_font("Helvetica", "", 9)
        col = epw / 2
        yb = pdf.get_y()
        pdf.set_xy(pdf.l_margin + col, yb)
        pdf.multi_cell(col, 5, latin(f"{_fmt_tgl(iso())}\nMengetahui,\nKepala {puskesmas}"), align="C")
        yb2 = pdf.get_y() + 20
        pdf.set_xy(pdf.l_margin + col, yb2)
        pdf.cell(col, 5, latin("( ................................ )"), align="C")
        out = io.BytesIO(pdf.output()); out.seek(0)
        return StreamingResponse(out, media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={fname}.pdf"})

    raise HTTPException(400, "Format tidak dikenal")

# ==================== KADER MANAGEMENT ====================
@api.get("/admin/kader")
async def list_kader(user=Depends(require_admin)):
    return await db.users.find({"role": "kader"}, {"_id": 0, "password_hash": 0}).to_list(200)

@api.get("/dashboard/rekap-kader")
async def rekap_kader(start: str = "", end: str = "", kelurahan: str = "", user=Depends(require_admin)):
    """Rekap per kader (tidak digabung): jumlah keluarga unik yang sudah ditangani tiap kader."""
    vq = {"status": "terkirim"}
    if kelurahan:
        vq["kelurahan"] = kelurahan
    if start or end:
        vq["created_at"] = dt_range(start, end)
    agg = {}
    pipeline = [
        {"$match": vq},
        {"$group": {"_id": "$kader_id",
                    "keluarga": {"$addToSet": "$keluarga_id"},
                    "total_kunjungan": {"$sum": 1}}},
    ]
    async for row in db.kunjungan.aggregate(pipeline):
        agg[row["_id"]] = {
            "keluarga_ditangani": len(row.get("keluarga", [])),
            "total_kunjungan": row.get("total_kunjungan", 0),
        }
    kaders = await db.users.find({"role": "kader"}, {"_id": 0, "password_hash": 0}).to_list(500)
    hasil = []
    for k in kaders:
        a = agg.get(k["id"], {"keluarga_ditangani": 0, "total_kunjungan": 0})
        draf = await db.kunjungan.count_documents({"kader_id": k["id"], "status": "draf"})
        target = k.get("target_keluarga", 0) or 0
        ditangani = a["keluarga_ditangani"]
        hasil.append({
            "kader_id": k["id"], "nama": k["nama"], "username": k["username"],
            "wilayah": k.get("wilayah", []), "posyandu": k.get("posyandu", ""),
            "target_keluarga": target, "aktif": k.get("aktif", True),
            "keluarga_ditangani": ditangani, "total_kunjungan": a["total_kunjungan"],
            "draf": draf,
            "cakupan": round(ditangani / target * 100, 1) if target else 0,
        })
    hasil.sort(key=lambda x: x["keluarga_ditangani"], reverse=True)
    total_unik = len(await db.kunjungan.distinct("keluarga_id", vq))
    return {"kader": hasil, "total_kader": len(hasil), "total_keluarga_ditangani_unik": total_unik}

@api.post("/admin/kader")
async def add_kader(body: dict, user=Depends(require_admin)):
    if not body.get("username") or not body.get("nama"):
        raise HTTPException(400, "Username dan nama wajib diisi")
    if await db.users.find_one({"username": body["username"].lower()}):
        raise HTTPException(400, "Username sudah dipakai")
    doc = {"id": new_id(), "username": body["username"].lower(), "nama": body["nama"],
           "role": "kader", "wilayah": body.get("wilayah", []), "posyandu": body.get("posyandu", ""),
           "target_keluarga": body.get("target_keluarga", 30), "aktif": True,
           "password_hash": hash_password(body.get("password", "kader123")), "created_at": iso()}
    await db.users.insert_one(doc); await audit(user, "create", "kader", doc["id"], body["nama"])
    doc.pop("_id", None); doc.pop("password_hash", None); return doc

@api.put("/admin/kader/{uid}")
async def update_kader(uid: str, body: dict, user=Depends(require_admin)):
    upd = {k: v for k, v in body.items() if k in ("nama", "wilayah", "posyandu", "target_keluarga", "aktif")}
    if body.get("password"):
        upd["password_hash"] = hash_password(body["password"])
    await db.users.update_one({"id": uid}, {"$set": upd})
    await audit(user, "update", "kader", uid, body.get("nama", ""))
    return await db.users.find_one({"id": uid}, {"_id": 0, "password_hash": 0})

class ResetPwIn(BaseModel):
    new_password: Optional[str] = None

@api.post("/admin/kader/{uid}/reset-password")
async def reset_kader_password(uid: str, body: ResetPwIn, user=Depends(require_admin)):
    u = await db.users.find_one({"id": uid, "role": "kader"})
    if not u:
        raise HTTPException(404, "Kader tidak ditemukan")
    newpw = (body.new_password or "").strip() or "kader123"
    if len(newpw) < 6:
        raise HTTPException(400, "Kata sandi minimal 6 karakter")
    await db.users.update_one({"id": uid}, {"$set": {"password_hash": hash_password(newpw)}})
    await audit(user, "reset_password", "kader", uid, u.get("nama", ""))
    return {"ok": True, "username": u["username"], "new_password": newpw}

@api.delete("/admin/kader/{uid}")
async def delete_kader(uid: str, user=Depends(require_admin)):
    u = await db.users.find_one({"id": uid, "role": "kader"})
    if not u:
        raise HTTPException(404, "Kader tidak ditemukan")
    await db.users.delete_one({"id": uid})
    await audit(user, "delete", "kader", uid, u.get("nama", ""))
    return {"ok": True, "nama": u.get("nama", "")}

# ==================== IMPORT / TEMPLATE EXCEL (ADMIN) ====================
KELURAHAN_LIST = ["Selat Tengah", "Selat Hulu", "Selat Dalam", "Selat Utara"]
KELUARGA_COLS = ["nama_kk", "kelurahan", "alamat", "rt", "rw", "posyandu", "no_hp", "catatan_lokasi", "punya_jkn", "air_bersih", "jamban", "ventilasi"]
KADER_COLS = ["nama", "username", "password", "wilayah", "posyandu", "target_keluarga"]

def _parse_bool_cell(v):
    s = str(v).strip().lower()
    if s in ("ya", "yes", "true", "1", "y"): return True
    if s in ("tidak", "no", "false", "0", "n"): return False
    return None

def _xlsx_response(wb, filename):
    out = io.BytesIO(); wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})

def _build_template(cols, example, panduan, title):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Data"
    for i, h in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D9488")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    ws.append(example)
    ws.freeze_panes = "A2"
    pan = wb.create_sheet("Panduan")
    pan.append([title]); pan["A1"].font = Font(bold=True, size=14, color="0D9488")
    pan.append([""])
    for line in panduan:
        pan.append([line])
    pan.column_dimensions["A"].width = 95
    return wb

@api.get("/admin/import/template/keluarga")
async def template_keluarga(user=Depends(require_admin)):
    panduan = [
        "1. Isi data mulai baris ke-2 pada sheet 'Data'. JANGAN mengubah nama kolom di baris 1.",
        "2. Kolom WAJIB: nama_kk, kelurahan.",
        f"3. Kolom 'kelurahan' harus salah satu dari: {', '.join(KELURAHAN_LIST)}.",
        "4. Kolom punya_jkn, air_bersih, jamban, ventilasi diisi 'Ya' atau 'Tidak' (boleh dikosongkan).",
        "5. Kolom rt, rw, no_hp diisi sebagai teks (mis. 01, 02).",
        "6. Baris dengan ALAMAT yang sama persis akan otomatis digabung menjadi 1 keluarga.",
        "7. Baris contoh (Budi Santoso) boleh dihapus sebelum diunggah.",
        "8. Simpan sebagai file .xlsx lalu unggah pada menu Import Data > Keluarga.",
    ]
    example = ["Budi Santoso", "Selat Tengah", "Jl. Melati No. 1", "01", "02", "Melati 1", "081234567890", "Dekat masjid", "Ya", "Ya", "Ya", "Ya"]
    wb = _build_template(KELUARGA_COLS, example, panduan, "PANDUAN IMPOR DATA KELUARGA")
    return _xlsx_response(wb, "template_keluarga.xlsx")

@api.get("/admin/import/template/kader")
async def template_kader(user=Depends(require_admin)):
    panduan = [
        "1. Isi data mulai baris ke-2 pada sheet 'Data'. JANGAN mengubah nama kolom di baris 1.",
        "2. Kolom WAJIB: nama, username, wilayah.",
        "3. 'username' harus unik (belum dipakai). Ditulis tanpa spasi, mis. kader11.",
        "4. 'password' opsional — jika dikosongkan, kader diberi kata sandi default 'kader123'.",
        f"5. 'wilayah' diisi satu kelurahan: {', '.join(KELURAHAN_LIST)}.",
        "6. 'target_keluarga' diisi angka (mis. 30). Jika kosong, otomatis 30.",
        "7. Baris contoh boleh dihapus sebelum diunggah. Simpan sebagai .xlsx lalu unggah pada menu Import Data > Kader.",
    ]
    example = ["Siti Aminah", "kader11", "kader123", "Selat Tengah", "Melati 1", 30]
    wb = _build_template(KADER_COLS, example, panduan, "PANDUAN IMPOR DATA KADER")
    return _xlsx_response(wb, "template_kader.xlsx")

def _read_sheet(content):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "File tidak valid. Gunakan file .xlsx dari template.")
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "File kosong.")
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    return header, rows[1:]

@api.post("/admin/import/keluarga")
async def import_keluarga(file: UploadFile = File(...), kader_id: Optional[str] = Form(default=None), user=Depends(require_admin)):
    header, data_rows = _read_sheet(await file.read())
    kader = None
    if kader_id:
        kader = await db.users.find_one({"id": kader_id, "role": "kader"})
        if not kader:
            raise HTTPException(400, "Kader tujuan tidak ditemukan")
    owner_id = kader["id"] if kader else user["id"]
    owner_nama = kader["nama"] if kader else user["nama"]
    created, digabung, errors = 0, 0, []
    seen = {}  # alamat_norm -> True (dalam file yang sama)
    for idx, raw in enumerate(data_rows, start=2):
        rec = {header[i]: raw[i] for i in range(len(header)) if i < len(raw)}
        nama = str(rec.get("nama_kk") or "").strip()
        kel = str(rec.get("kelurahan") or "").strip()
        if not nama and not kel and not any(v not in (None, "") for v in (raw or [])):
            continue
        if not nama or not kel:
            errors.append({"row": idx, "msg": "nama_kk dan kelurahan wajib diisi"}); continue
        alamat = str(rec.get("alamat") or "").strip()
        alamat_norm = " ".join(alamat.lower().split())
        if alamat_norm:
            # alamat yang sama persis digabung menjadi 1 keluarga
            if alamat_norm in seen:
                digabung += 1
                await audit(user, "import_gabung", "keluarga", "", f"baris {idx} digabung (alamat sama): {alamat}")
                continue
            existing = await db.keluarga.find_one({
                "deleted": {"$ne": True},
                "$or": [{"alamat_norm": alamat_norm},
                        {"alamat": {"$regex": f"^{re.escape(alamat)}$", "$options": "i"}}],
            })
            if existing:
                seen[alamat_norm] = True
                digabung += 1
                # jika ditujukan ke kader tertentu, alihkan kepemilikan keluarga yang sudah ada
                if kader:
                    await db.keluarga.update_one({"id": existing["id"]}, {"$set": {
                        "kader_id": kader["id"], "kader_nama": kader["nama"],
                        "created_by": kader["id"], "created_by_nama": kader["nama"]}})
                await audit(user, "import_gabung", "keluarga", existing["id"], f"baris {idx} digabung (alamat sama): {alamat}")
                continue
            seen[alamat_norm] = True
        doc = {
            "id": new_id(), "deleted": False, "created_by": owner_id, "created_by_nama": owner_nama,
            "kader_id": kader["id"] if kader else None, "kader_nama": kader["nama"] if kader else None,
            "created_at": iso(), "puskesmas": "UPT Puskesmas Melati", "data_lengkap": True,
            "nama_kk": nama, "kelurahan": kel,
            "alamat": alamat, "alamat_norm": alamat_norm,
            "rt": str(rec.get("rt") or "").strip(), "rw": str(rec.get("rw") or "").strip(),
            "posyandu": str(rec.get("posyandu") or "").strip() or (kader.get("posyandu", "") if kader else ""),
            "no_hp": str(rec.get("no_hp") or "").strip(),
            "catatan_lokasi": str(rec.get("catatan_lokasi") or "").strip(),
            "punya_jkn": _parse_bool_cell(rec.get("punya_jkn")),
            "air_bersih": _parse_bool_cell(rec.get("air_bersih")),
            "jamban": _parse_bool_cell(rec.get("jamban")),
            "ventilasi": _parse_bool_cell(rec.get("ventilasi")),
        }
        await db.keluarga.insert_one(doc); created += 1
    tujuan = f" ditujukan ke kader {kader['nama']}" if kader else ""
    await audit(user, "import", "keluarga", "", f"{created} dibuat, {digabung} digabung, {len(errors)} gagal{tujuan}")
    return {"ok": True, "created": created, "digabung": digabung, "gagal": len(errors),
            "errors": errors[:50], "kader": kader["nama"] if kader else None}

@api.post("/admin/import/kader")
async def import_kader(file: UploadFile = File(...), user=Depends(require_admin)):
    header, data_rows = _read_sheet(await file.read())
    created, errors = 0, []
    for idx, raw in enumerate(data_rows, start=2):
        rec = {header[i]: raw[i] for i in range(len(header)) if i < len(raw)}
        nama = str(rec.get("nama") or "").strip()
        username = str(rec.get("username") or "").strip().lower()
        wil = str(rec.get("wilayah") or "").strip()
        if not nama and not username and not any(v not in (None, "") for v in (raw or [])):
            continue
        if not nama or not username:
            errors.append({"row": idx, "msg": "nama dan username wajib diisi"}); continue
        if not wil:
            errors.append({"row": idx, "msg": "wilayah (kelurahan) wajib diisi"}); continue
        if await db.users.find_one({"username": username}):
            errors.append({"row": idx, "msg": f"username '{username}' sudah dipakai"}); continue
        pw = str(rec.get("password") or "").strip() or "kader123"
        try:
            target = int(rec.get("target_keluarga") or 30)
        except Exception:
            target = 30
        doc = {"id": new_id(), "username": username, "nama": nama, "role": "kader",
               "wilayah": [wil], "posyandu": str(rec.get("posyandu") or "").strip(),
               "target_keluarga": target, "aktif": True,
               "password_hash": hash_password(pw), "created_at": iso()}
        await db.users.insert_one(doc); created += 1
    await audit(user, "import", "kader", "", f"{created} dibuat, {len(errors)} gagal")
    return {"ok": True, "created": created, "gagal": len(errors), "errors": errors[:50]}

PERTANYAAN_COLS = ["group", "section", "text", "jenis", "satuan", "opsi", "wajib", "priority", "problem_when", "report_required", "definisi"]
_GROUP_LABELS = {k: v for k, v in KELOMPOK_LABEL.items() if k != "belum_ditentukan"}

@api.get("/admin/import/template/pertanyaan")
async def template_pertanyaan(user=Depends(require_admin)):
    grp = ", ".join(_GROUP_LABELS.keys())
    panduan = [
        "1. Isi data mulai baris ke-2 pada sheet 'Data'. JANGAN mengubah nama kolom di baris 1.",
        "2. Kolom WAJIB: group dan text.",
        f"3. 'group' harus salah satu kode: {grp}.",
        "4. 'section' diisi 'ceklis' (pertanyaan biasa) atau 'tanda_bahaya'. Kosong = ceklis.",
        "5. 'jenis' diisi salah satu: yesno, single, number, pemeriksaan, imunisasi, danger. Kosong = yesno.",
        "6. 'opsi' hanya untuk jenis 'single' — pisahkan pilihan dengan tanda titik-koma (;). Contoh: Ya; Tidak; Tidak tahu.",
        "7. 'satuan' untuk jenis 'number' (mis. °C, kg, cm).",
        "8. 'wajib' & 'report_required' diisi 'Ya' atau 'Tidak'.",
        "9. 'priority' diisi 'kuning' atau 'merah' (boleh dikosongkan).",
        "10. 'problem_when' = jawaban penanda masalah, pisahkan dengan titik-koma (;). Contoh: Tidak; Tidak tahu.",
        "11. Kode pertanyaan dibuat otomatis. Baris contoh boleh dihapus sebelum diunggah.",
    ]
    example = ["ibu_hamil", "ceklis", "Apakah ibu rutin minum tablet tambah darah?", "single", "", "Ya; Kadang; Tidak", "Ya", "kuning", "Tidak; Kadang", "Tidak", "Kepatuhan konsumsi TTD selama kehamilan"]
    wb = _build_template(PERTANYAAN_COLS, example, panduan, "PANDUAN IMPOR MASTER PERTANYAAN")
    return _xlsx_response(wb, "template_pertanyaan.xlsx")

def _split_multi(v):
    s = str(v or "").strip()
    if not s:
        return []
    parts = re.split(r"[;\n]", s) if (";" in s or "\n" in s) else s.split(",")
    return [p.strip() for p in parts if p.strip()]

@api.post("/admin/import/pertanyaan")
async def import_pertanyaan(file: UploadFile = File(...), user=Depends(require_admin)):
    header, data_rows = _read_sheet(await file.read())
    created, errors = 0, []
    urutan_cache = {}
    for idx, raw in enumerate(data_rows, start=2):
        rec = {header[i]: raw[i] for i in range(len(header)) if i < len(raw)}
        text = str(rec.get("text") or "").strip()
        group = str(rec.get("group") or "").strip().lower()
        if not text and not group and not any(v not in (None, "") for v in (raw or [])):
            continue
        if not text or not group:
            errors.append({"row": idx, "msg": "group dan text wajib diisi"}); continue
        if group not in _GROUP_LABELS:
            errors.append({"row": idx, "msg": f"group '{group}' tidak valid"}); continue
        section = str(rec.get("section") or "ceklis").strip().lower()
        if section not in QUESTION_SECTIONS:
            section = "ceklis"
        jenis = str(rec.get("jenis") or "yesno").strip().lower()
        if jenis not in QUESTION_JENIS:
            jenis = "yesno"
        base = re.sub(r"[^A-Z0-9]+", "_", f"{group}_{text[:20]}".upper()).strip("_") or group.upper()
        kode = base; n = 1
        while await db.master_questions.find_one({"kode": kode}):
            n += 1; kode = f"{base}_{n}"
        if group not in urutan_cache:
            last = await db.master_questions.find_one({"group": group}, sort=[("urutan", -1)])
            urutan_cache[group] = (last.get("urutan", 0) if last else 0)
        urutan_cache[group] += 1
        priority = str(rec.get("priority") or "").strip().lower() or None
        if priority not in (None, "kuning", "merah"):
            priority = None
        doc = {
            "kode": kode, "group": group, "section": section, "text": text,
            "definisi": str(rec.get("definisi") or "").strip(),
            "jenis": jenis, "satuan": (str(rec.get("satuan")).strip() or None) if rec.get("satuan") else None,
            "opsi": _split_multi(rec.get("opsi")), "wajib": _parse_bool_cell(rec.get("wajib")) or False,
            "problem_when": _split_multi(rec.get("problem_when")), "priority": priority,
            "report_required": _parse_bool_cell(rec.get("report_required")) or False,
            "urutan": urutan_cache[group], "deleted": False, "custom": True,
        }
        await db.master_questions.insert_one(doc); created += 1
    await audit(user, "import", "master_pertanyaan", "", f"{created} dibuat, {len(errors)} gagal")
    return {"ok": True, "created": created, "gagal": len(errors), "errors": errors[:50]}

# ==================== AUDIT LOG ====================
@api.get("/audit")
async def get_audit(user=Depends(require_admin)):
    return await db.audit_logs.find({}, {"_id": 0}).sort("waktu", -1).limit(200).to_list(200)

# ==================== AKREDITASI DASHBOARD ====================
@api.get("/akreditasi")
async def list_akreditasi(user=Depends(require_admin)):
    rows = await db.akreditasi.find({}, {"_id": 0}).sort("updated_at", -1).to_list(50)
    return rows

@api.get("/akreditasi/{did}")
async def get_akreditasi(did: str, user=Depends(require_admin)):
    d = await db.akreditasi.find_one({"id": did}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Dashboard tidak ditemukan")
    return d

@api.post("/akreditasi")
async def create_akreditasi(body: dict, user=Depends(require_admin)):
    from seed_demo import default_akreditasi
    doc = default_akreditasi()
    doc.update({k: v for k, v in body.items() if k not in ("id",)})
    doc["id"] = new_id(); doc["created_at"] = iso(); doc["updated_at"] = iso()
    doc["updated_by"] = user["nama"]; doc.setdefault("versions", [])
    await db.akreditasi.insert_one(doc); await audit(user, "create", "akreditasi", doc["id"], doc.get("judul", ""))
    doc.pop("_id", None); return doc

@api.put("/akreditasi/{did}")
async def update_akreditasi(did: str, body: dict, user=Depends(require_admin)):
    cur = await db.akreditasi.find_one({"id": did}, {"_id": 0})
    if not cur:
        raise HTTPException(404, "Tidak ditemukan")
    versions = cur.get("versions", [])
    snap = {k: v for k, v in cur.items() if k != "versions"}
    versions.append({"waktu": iso(), "oleh": user["nama"], "snapshot": snap})
    versions = versions[-20:]
    body.pop("_id", None); body.pop("id", None); body.pop("versions", None)
    body["updated_at"] = iso(); body["updated_by"] = user["nama"]; body["versions"] = versions
    await db.akreditasi.update_one({"id": did}, {"$set": body})
    await audit(user, "update", "akreditasi", did, body.get("judul", ""))
    return await db.akreditasi.find_one({"id": did}, {"_id": 0})

@api.post("/akreditasi/{did}/restore/{idx}")
async def restore_akreditasi(did: str, idx: int, user=Depends(require_admin)):
    cur = await db.akreditasi.find_one({"id": did}, {"_id": 0})
    if not cur or idx >= len(cur.get("versions", [])):
        raise HTTPException(404, "Versi tidak ditemukan")
    snap = cur["versions"][idx]["snapshot"]
    snap["versions"] = cur["versions"]; snap["updated_at"] = iso()
    await db.akreditasi.update_one({"id": did}, {"$set": snap})
    return await db.akreditasi.find_one({"id": did}, {"_id": 0})

@api.delete("/akreditasi/{did}")
async def delete_akreditasi(did: str, user=Depends(require_admin)):
    await db.akreditasi.delete_one({"id": did}); return {"ok": True}

app.include_router(api)
_frontend_url = os.environ.get("FRONTEND_URL", "").strip()
_cors_env = [o.strip() for o in os.environ.get("CORS_ORIGINS", "").split(",") if o.strip() and o.strip() != "*"]
_allowed_origins = list({o for o in ([_frontend_url] if _frontend_url else []) + _cors_env})
app.add_middleware(CORSMiddleware, allow_credentials=True,
                   allow_origins=_allowed_origins,
                   allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?|https://[a-z0-9.-]+\.(emergentagent\.com|emergentcf\.cloud)",
                   allow_methods=["*"], allow_headers=["*"])

@app.on_event("shutdown")
async def shutdown():
    client.close()
